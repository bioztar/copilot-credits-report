"""Extract every AI cost / usage source the signed-in admin can reach, into one Excel workbook.

For each source it records what was called, whether it worked, and (if not) which role is missing.
Output, all kept: <usage-dir>/AI_Cost_Extraction/<date_time>/
    AI_Cost_Sources_<date_time>.xlsx   Access (status + role to ask for), Queries (every HTTP call), one tab per data set
    raw/NNN_<source>.json              every response exactly as returned
    run.log                            console output

Sign-in is delegated, read-only, device code. Two clients: Microsoft Graph Command Line Tools (Copilot cost
management, reuses pull_credits' sign-in) and Azure CLI (Azure, Power Platform, Graph directory). Tokens stay in
~/.config/copilot-credits/ and are never printed or saved in the output.

    python extract_ai_costs.py "<path to Usage folder>"
"""
import datetime as dt, io, json, pathlib, re, sys, time, traceback
import msal, pandas as pd, requests
from openpyxl.styles import Font, PatternFill
import pull_credits as pc

AZ_CLIENT = "04b07795-8ddb-461a-bbee-02f9e1bf7b46"  # Azure CLI
AZ_CACHE = pathlib.Path.home() / ".config/copilot-credits/msal_cache_azcli.json"
ARM, PPAPI, G = "https://management.azure.com", "https://api.powerplatform.com", "https://graph.microsoft.com"
AI = re.compile(r"copilot|virtual.?agent|ccibots|power.?virtual|ai.?builder|agent|credit|openai|cognitive|"
                r"machinelearning|foundry|powerplatform|power platform", re.I)

# Where to ask when a source is out of reach.
ROLE = {
    "graph_cost": ("Entra admin role + delegated CopilotCostManagement-* permissions", "Global Admin (already granted to Vitaly)"),
    "graph_usage": ("Reports Reader (or higher) + Reports.Read.All", "Global Admin"),
    "graph_skus": ("Any directory read role (Global Reader / License Administrator)", "Global Admin"),
    "audit": ("Purview Audit Reader role + AuditLogsQuery.Read.All consent", "Security (Roxanne Andrews), plan task #224"),
    "agents": ("Not a role: the API needs an Agent 365 licence in the tenant (then CopilotPackages.Read.All)", "Licensing / Microsoft account team"),
    "billing": ("MCA: Billing account reader / EA: Enterprise read-only", "Finance / owner of the Microsoft billing account"),
    "azure_cost": ("Cost Management Reader on the subscription that pays the Copilot / Copilot Studio billing policies "
                   "(not visible to this user today). If even Contributor gets 401, the Enterprise Agreement setting "
                   "'Account owner can view charges' is off and only the EA admin can switch it on", "Owner of that Azure subscription (IT cloud team)"),
    "azure_res": ("Reader on the subscription", "Owner of the IT Azure subscription"),
    "pp": ("Power Platform Administrator (Entra role)", "Global Admin"),
    "github": ("GitHub enterprise owner or billing manager", "GitHub enterprise owner"),
}


class AzSignIn:
    """One Azure CLI device-code sign-in, then silent tokens for any resource that client may call."""

    def __init__(self, say):
        self.say, self.cache = say, msal.SerializableTokenCache()
        if AZ_CACHE.exists():
            self.cache.deserialize(AZ_CACHE.read_text())
        self.app = msal.PublicClientApplication(AZ_CLIENT, token_cache=self.cache,
                                                authority="https://login.microsoftonline.com/organizations")

    def token(self, resource: str) -> str:
        scopes = [f"{resource}/.default"]
        accts = self.app.get_accounts()
        res = self.app.acquire_token_silent(scopes, account=accts[0]) if accts else None
        if not res:
            flow = self.app.initiate_device_flow(scopes=scopes)
            self.say(flow.get("message", str(flow)))
            res = self.app.acquire_token_by_device_flow(flow)
        if self.cache.has_state_changed:
            AZ_CACHE.parent.mkdir(parents=True, exist_ok=True)
            AZ_CACHE.write_text(self.cache.serialize())
            AZ_CACHE.chmod(0o600)
        if "access_token" not in res:
            raise RuntimeError(f"sign-in for {resource} failed: {res.get('error')}: "
                               f"{(res.get('error_description') or '')[:200]}")
        return res["access_token"]


def err_text(r: requests.Response) -> str:
    try:
        e = r.json().get("error", r.json())
        msg = e.get("message") or e.get("code") or e if isinstance(e, dict) else e
    except ValueError:
        msg = r.text
    return f"HTTP {r.status_code} {r.reason}: {str(msg)[:300]}".strip()


class Run:
    def __init__(self, out: pathlib.Path):
        self.out, self.raw = out, out / "raw"
        self.raw.mkdir(parents=True, exist_ok=True)
        self.logf = open(out / "run.log", "a", encoding="utf-8")
        self.queries, self.access, self.tabs = [], [], {}

    def say(self, msg: str):
        print(msg, flush=True)
        self.logf.write(msg + "\n"); self.logf.flush()

    def call(self, source: str, method: str, url: str, tok: str, body=None) -> requests.Response:
        for _ in range(6):
            r = requests.request(method, url, headers={"Authorization": f"Bearer {tok}"}, json=body, timeout=120)
            if r.status_code != 429:
                break
            time.sleep(int(r.headers.get("Retry-After", 10)))
        n = len(self.queries) + 1
        fname = f"{n:03d}_{re.sub(r'[^a-z0-9]+', '_', source.lower()).strip('_')}.json"
        try:
            payload = r.json()
        except ValueError:
            payload = r.text
        (self.raw / fname).write_text(json.dumps({"method": method, "url": url, "body": body, "status": r.status_code,
                                                  "response": payload}, indent=1, ensure_ascii=False))
        self.queries.append({"#": n, "Source": source, "Method": method, "URL": url,
                             "Body": json.dumps(body) if body else "", "HTTP": r.status_code,
                             "Result": "OK" if r.ok else err_text(r), "Raw file": f"raw/{fname}",
                             "At (UTC)": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d %H:%M:%S")})
        return r

    def pages(self, source: str, url: str, tok: str):
        """Yield items across @odata.nextLink pages. Raises on the first failed page."""
        while url:
            r = self.call(source, "GET", url, tok)
            if not r.ok:
                raise PermissionError(err_text(r))
            body = r.json()
            yield from body.get("value", [])
            url = body.get("@odata.nextLink") or body.get("nextLink")
            if url:  # Graph bug: the ByDate report's nextLink drops "/copilot" and then 400s
                url = url.replace("/beta/reports/getMicrosoft365CopilotUsageUserDetailByDate",
                                  "/beta/copilot/reports/getMicrosoft365CopilotUsageUserDetailByDate")

    def record(self, source, model, endpoint, query, status, rows=0, detail="", role_key=None, tab=""):
        role, ask = ROLE.get(role_key, ("", ""))
        self.access.append({"Source": source, "Data model / what it holds": model, "API / endpoint": endpoint,
                            "Query run": query, "Status": status, "Rows": rows, "Tab": tab, "Detail": detail,
                            "Role / permission needed": role if status != "OK" else "",
                            "Ask": ask if status != "OK" else ""})
        self.say(f"[{status}] {source}: {rows} rows {detail}")

    def tab(self, name: str, rows) -> str:
        df = rows if isinstance(rows, pd.DataFrame) else pd.json_normalize(rows) if rows else pd.DataFrame()
        df = df[[c for c in df.columns if "@odata" not in c]]
        self.tabs[name[:31]] = df
        return name[:31]


def attempt(run: Run, source, model, endpoint, query, role_key, fn):
    """Run one extraction; fn returns (rows, tab_name, detail). Any failure is recorded, never fatal."""
    try:
        rows, tab, detail = fn()
        run.record(source, model, endpoint, query, "OK" if rows else "OK (empty)", len(rows), detail, role_key, tab)
    except PermissionError as e:
        run.record(source, model, endpoint, query, "NO ACCESS", 0, str(e), role_key)
    except (Exception, SystemExit) as e:  # one source failing must not stop the others
        run.logf.write(traceback.format_exc())
        run.record(source, model, endpoint, query, "ERROR", 0, f"{type(e).__name__}: {str(e)[:300]}", role_key)


def get_json(run, source, url, tok):
    r = run.call(source, "GET", url, tok)
    if not r.ok:
        raise PermissionError(err_text(r))
    return r.json()


# ---------- Graph: Copilot cost management (credits) ----------

def graph_cost(run: Run):
    tok = pc.token(True)
    src = "Graph Copilot Cost Management"
    state = {}

    def balance():
        rows = get_json(run, src, f"{pc.CM}/getTenantCreditBalance()", tok).get("billingMethodBalances", [])
        return rows, run.tab("Graph_TenantBalance", rows), "credits per billing method (PAYG, prepaid)"

    def policies():
        state["pol"] = list(run.pages(src, f"{pc.CM}/spendingPolicies", tok))
        return state["pol"], run.tab("Graph_SpendingPolicies", state["pol"]), ""

    def settings():
        rows = [{"policy": p.get("displayName"), "policyId": p["id"], **s}
                for p in state.get("pol", []) for s in run.pages(src, f"{pc.CM}/spendingPolicies/{p['id']}/serviceSettings", tok)]
        return rows, run.tab("Graph_PolicyServices", rows), "per-policy service on/off and limits"

    def user_services():
        users = {}
        for p in state.get("pol", []):
            for g in run.pages(src, f"{pc.CM}/spendingPolicies/{p['id']}/microsoft.graph.selectedGroupsSpendingPolicy/assignedGroups", tok):
                for u in run.pages(src, f"{G}/v1.0/groups/{g['id']}/transitiveMembers/microsoft.graph.user"
                                        "?$select=id,userPrincipalName,displayName&$top=999", tok):
                    users[u["id"]] = {**u, "policy": p.get("displayName")}
        rows = [{"upn": u["userPrincipalName"], "displayName": u["displayName"], "policy": u["policy"], **b}
                for uid, u in users.items() for b in run.pages(src, f"{pc.CM}/userBalances/{uid}/serviceBalances", tok)]
        if rows:  # one row per user, one column per service: the view to read first
            piv = pd.DataFrame(rows).pivot_table(index=["upn", "displayName", "policy"], columns="serviceId",
                                                 values="consumedQuantity", aggfunc="sum", fill_value=0)
            piv = piv.loc[:, piv.sum() > 0].assign(total=lambda d: d.sum(axis=1)).sort_values("total", ascending=False)
            run.tab("Graph_CreditsPerUser", piv.round(2).reset_index())
        return rows, run.tab("Graph_UserServiceCredits", rows), f"{len(users)} users in policy groups, current month"

    ep = "graph.microsoft.com/beta/copilot/costManagement"
    attempt(run, src, "Tenant credit balance (credits, current period)", ep, "GET /getTenantCreditBalance()", "graph_cost", balance)
    attempt(run, src, "Spending policies (tiers, limits)", ep, "GET /spendingPolicies", "graph_cost", policies)
    attempt(run, src, "Services enabled per policy", ep, "GET /spendingPolicies/{id}/serviceSettings", "graph_cost", settings)
    attempt(run, src, "Credits per user x service (Cowork, Studio, Code, Work IQ...)", ep,
            "GET /spendingPolicies/{id}/…/assignedGroups → /groups/{id}/transitiveMembers → /userBalances/{userId}/serviceBalances",
            "graph_cost", user_services)


# ---------- Graph: Copilot usage reports, licences, audit, agents ----------

def graph_usage(run: Run):
    tok = pc.token(True)
    src = "Graph Copilot usage reports"

    def detail():
        for period in ("D180", "D90", "D28"):
            r = run.call(src, "GET", f"{G}/v1.0/copilot/reports/getMicrosoft365CopilotUsageUserDetail(period='{period}')", tok)
            if r.ok:
                df = pd.read_csv(io.StringIO(r.content.decode("utf-8-sig")))
                return df.to_dict("records"), run.tab("Graph_CopilotUsageUsers", df), f"period {period}"
        raise PermissionError(err_text(r))

    def summary(name):
        def f():
            rows = list(run.pages(src, f"{G}/beta/reports/getMicrosoft365Copilot{name}(period='D180')?$format=application/json", tok))
            return rows, run.tab(f"Graph_Copilot{name[9:]}", rows), "period D180"
        return f

    def by_date():
        # Microsoft only serves the 7 days before the report's refresh date (a day or two behind today):
        # try the last 10 days and keep the ones it accepts. History beyond 7 days = store each run's output.
        rows, last_err = [], None
        for back in range(1, 11):
            day = (dt.date.today() - dt.timedelta(days=back)).isoformat()
            try:
                rows += list(run.pages(src, f"{G}/beta/copilot/reports/getMicrosoft365CopilotUsageUserDetailByDate(date={day})"
                                            "?$format=application/json", tok))
            except PermissionError as e:
                last_err = e
        if not rows and last_err:
            raise last_err
        return rows, run.tab("Graph_CopilotUsageDaily", rows), "user x day, the 7 days Microsoft keeps"

    ep = "graph.microsoft.com/…/reports/getMicrosoft365Copilot*"
    attempt(run, src, "Licensed Copilot usage per user per day (new Oct 2026 function)", ep,
            "GET /beta/copilot/reports/getMicrosoft365CopilotUsageUserDetailByDate(date=…) for each of the last 10 days (7 served)",
            "graph_usage", by_date)
    attempt(run, src, "Licensed M365 Copilot activity per user (prompts, last use per app)", ep,
            "GET /v1.0/copilot/reports/getMicrosoft365CopilotUsageUserDetail(period='D180')", "graph_usage", detail)
    for name in ("UserCountSummary", "UserCountTrend"):
        attempt(run, src, f"Copilot {name}", ep, f"GET /beta/reports/getMicrosoft365Copilot{name}(period='D180')",
                "graph_usage", summary(name))


def graph_licences(run: Run, az: AzSignIn):
    src = "Graph licences"

    def skus():
        rows = [{"skuPartNumber": s["skuPartNumber"], "AI related": bool(AI.search(s["skuPartNumber"])),
                 "assigned": s["consumedUnits"], "bought (enabled)": s["prepaidUnits"]["enabled"],
                 "suspended": s["prepaidUnits"]["suspended"], "warning": s["prepaidUnits"]["warning"],
                 "capabilityStatus": s["capabilityStatus"], "skuId": s["skuId"],
                 "servicePlans": ", ".join(p["servicePlanName"] for p in s["servicePlans"])}
                for s in get_json(run, src, f"{G}/v1.0/subscribedSkus", az.token(G))["value"]]
        rows.sort(key=lambda x: (not x["AI related"], -x["bought (enabled)"]))
        return rows, run.tab("Graph_Licences", rows), "counts only, Graph has no prices"

    attempt(run, src, "Licences bought vs assigned per SKU (no prices)", "graph.microsoft.com/v1.0/subscribedSkus",
            "GET /v1.0/subscribedSkus", "graph_skus", skus)


def graph_extra(run: Run):
    """Sources that need an extra Graph consent; a refused consent is recorded as NO ACCESS."""
    checks = [
        ("Copilot agent catalog (Graph)", "Agents / Copilot packages in the tenant (inventory to join with costs)",
         "CopilotPackages.Read.All", f"{G}/beta/copilot/admin/catalog/packages", "agents", "Graph_AgentCatalog"),
    ]
    for src, model, scope, url, key, tab in checks:
        def f(src=src, scope=scope, url=url, tab=tab):
            tok = pc.token(True, [f"{G}/{scope}"])
            rows = list(run.pages(src, url, tok))
            return rows, run.tab(tab, rows), ""
        attempt(run, src, model, url.replace("https://", ""), f"GET {url.replace(G, '')}  (scope {scope})", key, f)


def audit_probe(run: Run, az: AzSignIn):
    """Listing audit searches works without the role; only creating one proves it. Creates one small search
    (signed-in user's own Copilot interactions, last 24 h) and leaves it; results are read on a later run."""
    src = "Purview audit log (Graph)"

    def f():
        me = get_json(run, src, f"{G}/v1.0/me?$select=userPrincipalName", az.token(G))["userPrincipalName"]
        upn = me.replace("_admin@", "@")  # admin accounts do not use Copilot; their day-to-day twin does
        now = dt.datetime.now(dt.timezone.utc)
        body = {"displayName": f"access test {now:%Y-%m-%d}", "recordTypeFilters": ["copilotInteraction"],
                "filterStartDateTime": f"{now - dt.timedelta(days=1):%Y-%m-%dT%H:%M:%SZ}",
                "filterEndDateTime": f"{now:%Y-%m-%dT%H:%M:%SZ}", "userPrincipalNameFilters": [upn]}
        r = run.call(src, "POST", f"{G}/v1.0/security/auditLog/queries", pc.token(True, [f"{G}/AuditLogsQuery.Read.All"]), body)
        if not r.ok:
            raise PermissionError(err_text(r))
        return [r.json()], run.tab("Graph_AuditSearch", [r.json()]), "search created; records readable once it succeeds"

    attempt(run, src, "Copilot / Cowork interaction audit records (who, when, which agent / app)",
            "graph.microsoft.com/v1.0/security/auditLog/queries", "POST one search: copilotInteraction, own user, 24 h",
            "audit", f)


def my_roles(run: Run, az: AzSignIn):
    src = "Signed-in user's Entra roles"

    def f():
        rows = list(run.pages(src, f"{G}/v1.0/me/memberOf/microsoft.graph.directoryRole?$select=displayName,description", az.token(G)))
        return rows, run.tab("Me_EntraRoles", rows), ", ".join(r["displayName"] for r in rows)

    attempt(run, src, "What the signing-in account can already do", "graph.microsoft.com/v1.0/me/memberOf",
            "GET /me/memberOf/microsoft.graph.directoryRole", None, f)


# ---------- Azure: billing, cost management, AI resources ----------

def azure(run: Run, az: AzSignIn):
    tok, src = az.token(ARM), "Azure"
    state = {}

    def billing():
        rows = [{"name": a["name"], **a["properties"]} for a in
                get_json(run, src, f"{ARM}/providers/Microsoft.Billing/billingAccounts?api-version=2024-04-01", tok)["value"]]
        if not rows:
            raise PermissionError("no billing account visible to this user (EA vs MCA unknown)")
        return rows, run.tab("Azure_BillingAccounts", rows), "agreementType = EA or MCA"

    def subs():
        state["subs"] = get_json(run, src, f"{ARM}/subscriptions?api-version=2022-12-01", tok)["value"]
        return state["subs"], run.tab("Azure_Subscriptions", state["subs"]), ""

    def cost():
        start = (dt.date.today().replace(day=1) - dt.timedelta(days=62)).replace(day=1).isoformat()
        body = {"type": "ActualCost", "timeframe": "Custom",
                "timePeriod": {"from": f"{start}T00:00:00Z", "to": f"{dt.date.today().isoformat()}T23:59:59Z"},
                "dataset": {"granularity": "Monthly", "aggregation": {"cost": {"name": "Cost", "function": "Sum"}},
                            "grouping": [{"type": "Dimension", "name": "MeterCategory"},
                                         {"type": "Dimension", "name": "Meter"}]}}
        rows, denied = [], []
        for s in state.get("subs", []):
            r = run.call(src, "POST", f"{ARM}/subscriptions/{s['subscriptionId']}/providers/Microsoft.CostManagement/query"
                                      "?api-version=2023-11-01", tok, body)
            if not r.ok:
                denied.append(f"{s['displayName']}: {err_text(r)}")
                continue
            p = r.json()["properties"]
            cols = [c["name"] for c in p["columns"]]
            rows += [{"subscription": s["displayName"], **dict(zip(cols, x))} for x in p["rows"]]
        for x in rows:
            x["AI related"] = bool(AI.search(f"{x.get('MeterCategory')} {x.get('Meter')}"))
        if denied and not rows:
            raise PermissionError("; ".join(denied))
        return rows, run.tab("Azure_CostByMeter", rows), "; ".join(denied)

    def resources():
        rows = [{"subscription": s["displayName"], "name": x["name"], "type": x["type"], "location": x.get("location"),
                 "resourceGroup": x["id"].split("/")[4], "sku": (x.get("sku") or {}).get("name"), "id": x["id"]}
                for s in state.get("subs", [])
                for x in run.pages(src, f"{ARM}/subscriptions/{s['subscriptionId']}/resources?api-version=2021-04-01", tok)
                if AI.search(x["type"])]
        return rows, run.tab("Azure_AIResources", rows), "Copilot billing policies, OpenAI/Foundry, ML"

    attempt(run, src, "Billing accounts (EA or MCA, licence invoices)", "management.azure.com/providers/Microsoft.Billing/billingAccounts",
            "GET billingAccounts?api-version=2024-04-01", "billing", billing)
    attempt(run, src, "Subscriptions visible", "management.azure.com/subscriptions", "GET subscriptions?api-version=2022-12-01",
            "azure_res", subs)
    attempt(run, src, "Actual cost per month x meter (EUR, the invoiced money)",
            "management.azure.com/subscriptions/{id}/providers/Microsoft.CostManagement/query",
            "POST query: ActualCost, last 3 months, Monthly, group by MeterCategory + Meter", "azure_cost", cost)
    attempt(run, src, "AI resources (billing policies, OpenAI / Foundry, ML)", "management.azure.com/subscriptions/{id}/resources",
            "GET resources, filtered locally to AI resource types", "azure_res", resources)


# ---------- Power Platform: Copilot Studio credits per agent ----------

def power_platform(run: Run, az: AzSignIn):
    tok, src = az.token(PPAPI), "Power Platform API"
    frm = (dt.date.today() - dt.timedelta(days=30)).isoformat()
    calls = [
        ("Copilot Studio credit capacity (tenant)", "/licensing/entitlements/MCSMessages?api-version=2024-10-01", "PP_StudioCapacity"),
        ("Copilot Studio credits per agent x day x environment (last 30 days)",
         f"/licensing/entitlements/MCSMessages/resources?fromDate={frm}&toDate={dt.date.today().isoformat()}"
         "&includeFields=users%2Ctags%2CasOfDate&pageSize=5000&api-version=2024-10-01", "PP_StudioAgentCredits"),
        ("Billing policies (link environments to the Azure subscription)",
         "/licensing/billingPolicies?api-version=2022-03-01-preview", "PP_BillingPolicies"),
        ("Environments (names for the agent rows)", "/environmentmanagement/environments?api-version=2024-10-01", "PP_Environments"),
    ]
    for model, path, tab in calls:
        def f(path=path, tab=tab):
            body = get_json(run, src, f"{PPAPI}{path}", tok)
            rows = body.get("value", [body]) if isinstance(body, dict) else body
            return rows, run.tab(tab, rows), ""
        attempt(run, src, model, f"api.powerplatform.com{path.split('?')[0]}", f"GET {path}", "pp", f)


def pp_inventory(run: Run, az: AzSignIn):
    """Power Platform inventory: every Copilot Studio agent with model, channels, triggers, owner. Readable without
    the Power Platform admin role. A query returns at most 1000 rows, so it is split by first hex digit of the id."""
    src = "Power Platform inventory"

    def f():
        users = {u["id"]: u for u in run.pages(src, f"{G}/v1.0/users?$select=id,userPrincipalName,department&$top=999", az.token(G))}
        tok, rows = az.token(PPAPI), []
        for h in "0123456789abcdef":
            body = {"TableName": "PowerPlatformResources", "Clauses": [
                {"$type": "where", "FieldName": "type", "Operator": "==", "Values": ["'microsoft.copilotstudio/agents'"]},
                {"$type": "where", "FieldName": "name", "Operator": "startswith", "Values": [f"'{h}'"]},
                {"$type": "take", "TakeCount": 1000}]}
            r = run.call(src, "POST", f"{PPAPI}/resourcequery/resources/query?api-version=2024-10-01", tok, body)
            if not r.ok:
                raise PermissionError(err_text(r))
            rows += r.json().get("data", [])
        out = []
        for a in rows:
            p, owner = a.get("properties") or {}, users.get((a.get("properties") or {}).get("ownerId"), {})
            cc, sv = p.get("componentsCounts") or {}, p.get("sharedWithViewers") or {}
            out.append({"agent": p.get("displayName"), "agentId": a["name"], "environmentId": p.get("environmentId"),
                        "createdIn": p.get("createdIn"), "harness": p.get("harness"), "model": p.get("model"),
                        "orchestration": p.get("orchestration"), "channels": ", ".join(p.get("channels") or []),
                        "autonomous (has triggers)": bool(p.get("triggers")), "triggers": len(p.get("triggers") or []),
                        "flows": len(p.get("flows") or []), "connectors": len(p.get("powerPlatformConnectors") or []),
                        "knowledge": cc.get("knowledge"), "tools": cc.get("tools"), "connectedAgents": cc.get("connectedAgents"),
                        "webSearch": p.get("isWebSearchEnabledForKnowledge"), "sharedWithTenant": sv.get("entireTenant"),
                        "sharedUsers": sv.get("userCount"), "sharedGroups": sv.get("groupCount"),
                        "authentication": p.get("authentication"), "quarantined": p.get("isQuarantined"),
                        "createdAt": p.get("createdAt"), "lastPublishedAt": p.get("lastPublishedAt"),
                        "ownerUpn": owner.get("userPrincipalName"), "ownerDepartment": owner.get("department"),
                        "entraAgentId": p.get("entraAgentId")})
        return out, run.tab("PP_AgentInventory", out), f"{sum(x['autonomous (has triggers)'] for x in out)} autonomous"

    attempt(run, src, "Every Copilot Studio agent: model, channels, autonomous triggers, sharing, owner + department",
            "api.powerplatform.com/resourcequery/resources/query", "POST KQL: type == copilotstudio/agents, 16 id-prefix slices",
            "pp", f)


def dataverse(run: Run, az: AzSignIn):
    """Per environment: Copilot Studio agents (bots) and sessions per agent in the last 30 days (conversationtranscripts).
    Needs System Administrator / Bot Transcript Viewer in each environment; the Power Platform Administrator role gives it everywhere."""
    src = "Dataverse (per environment)"

    def f():
        envs = get_json(run, src, f"{PPAPI}/environmentmanagement/environments?api-version=2024-10-01", az.token(PPAPI))["value"]
        since = (dt.date.today() - dt.timedelta(days=30)).isoformat()
        rows, denied = [], 0
        inv = run.tabs.get("PP_AgentInventory", pd.DataFrame())
        hosts = set(inv["environmentId"]) if "environmentId" in inv else None  # only environments that hold agents
        for e in [e for e in envs if e.get("url") and (hosts is None or e["id"] in hosts)]:
            base, name = e["url"].rstrip("/") + "/api/data/v9.2", e.get("displayName")
            try:
                tok = az.token(e["url"].rstrip("/"))
            except RuntimeError:
                denied += 1
                continue
            probe = run.call(src, "GET", f"{base}/conversationtranscripts?$top=1&$select=createdon", tok)
            if not probe.ok:
                denied += 1
                continue
            r = run.call(src, "GET", f"{base}/conversationtranscripts?$apply=filter(createdon ge {since})/groupby((_bot_conversationtranscriptid_value),"
                                     "aggregate($count as sessions))", tok)
            if not r.ok:
                denied += 1
                continue
            bots = run.call(src, "GET", f"{base}/bots?$select=botid,name,createdon,publishedon", tok)
            names = {b["botid"]: b["name"] for b in (bots.json().get("value", []) if bots.ok else [])}
            rows += [{"environment": name, "environmentId": e["id"], "agentId": x["_bot_conversationtranscriptid_value"],
                      "agent": names.get(x["_bot_conversationtranscriptid_value"]), "sessions last 30 days": x["sessions"]}
                     for x in r.json().get("value", [])]
        if denied and not rows:
            raise PermissionError(f"no readable transcripts in any of {denied} environments")
        return rows, run.tab("DV_AgentSessions", rows), f"transcripts unreadable in {denied} environments"

    attempt(run, src, "Copilot Studio sessions per agent (conversation transcripts), all environments",
            "{environment}.crm4.dynamics.com/api/data/v9.2/conversationtranscripts + bots",
            "GET conversationtranscripts $apply groupby bot, last 30 days; GET bots", "pp", f)


# ---------- Service coverage: which Copilot services the Graph data covers ----------

# serviceId fragment -> (plain name, note). Matched case-insensitively against the policy's service list.
SERVICES = [
    ("cowork", "Copilot Cowork", "Per user via Graph serviceBalances."),
    ("copilotCode", "Copilot Code", "New Microsoft 365 Copilot capability in the Copilot stack, next to Cowork. "
                                    "NOT GitHub Copilot (that is gitHubCopilotApp). Per user via Graph serviceBalances."),
    ("gitHubCopilotApp", "GitHub Copilot app in Microsoft 365", "Enabled in the policy; GitHub Copilot seats and premium "
                                                                 "requests are billed in GitHub, not here."),
    ("autopilot", "Copilot Autopilot", "Announced, expected soon, billed in Copilot credits like Cowork. Will appear in "
                                       "the policy service list and serviceBalances with no extra access."),
    ("copilotStudio", "Copilot Studio agents", "Enabled in the policy but no per-user balance; per-agent credits only "
                                               "via the Power Platform API."),
    ("workIQ", "Work IQ API", "Per user via Graph serviceBalances."),
]


def service_coverage(run: Run):
    pol = run.tabs.get("Graph_PolicyServices", pd.DataFrame())
    bal = run.tabs.get("Graph_UserServiceCredits", pd.DataFrame())
    ids = list(pol["serviceId"]) if "serviceId" in pol else []
    for frag, name, note in SERVICES:
        hit = [i for i in ids if frag.lower() in i.lower()]
        used = bal[bal["serviceId"].isin(hit)]["consumedQuantity"].sum() if hit and "serviceId" in bal else 0
        status = "OK" if hit else "NOT YET IN TENANT"
        run.record(f"Service: {name}", note, "graph …/costManagement/spendingPolicies/{id}/serviceSettings + serviceBalances",
                   "from the Graph Copilot Cost Management calls above", status, len(hit),
                   f"serviceId {', '.join(hit) or frag + ' (absent)'}; credits this month {used:,.0f}",
                   "pp" if frag == "copilotStudio" else None)


def not_automated(run: Run):
    run.record("GitHub Copilot billing", "GitHub Copilot seats / premium requests (in TEC-05 chargeback scope)",
               "GET api.github.com/enterprises/{enterprise}/settings/billing/usage", "not run: needs a GitHub token",
               "NOT RUN", 0, "API exists, separate GitHub sign-in", "github")


# ---------- Workbook ----------

def write_workbook(run: Run, path: pathlib.Path, stamp: str):
    navy, white = PatternFill("solid", fgColor="002957"), Font(bold=True, color="FFFFFF")
    colour = {"OK": "C6EFCE", "OK (empty)": "FFF2CC", "NO ACCESS": "FFC7CE", "ERROR": "FFC7CE", "NOT RUN": "E7E6E6", "NOT YET IN TENANT": "FFF2CC"}
    readme = pd.DataFrame({"AI cost & usage source extraction": [
        f"Run: {stamp} (local time). Internal use.",
        "Access: every source tried, status, and for each one not reached the role to ask for and who grants it.",
        "Queries: every HTTP call made, method, URL, body, HTTP status, and the raw response file (raw/ folder next to this workbook).",
        "Other tabs: the data each source returned, flattened. 'AI related' columns flag Copilot / agent / AI rows.",
        "API sources only: everything here can be pulled unattended or on a schedule; no manually exported files are read.",
        "Graph returns credits (1 credit = 0.01 USD PAYG list price); only Azure Cost Management returns invoiced money.",
        "Contains named per-user data: keep this folder restricted."]})
    sheets = {"README": readme, "Access": pd.DataFrame(run.access), "Queries": pd.DataFrame(run.queries), **run.tabs}
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            for c in ws[1]:
                c.fill, c.font = navy, white
            ws.freeze_panes, ws.auto_filter.ref = "A2", ws.dimensions
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(60, max(10, *(len(str(c.value or "")) for c in col[:200])) + 2)
            if name == "Access":
                st = [c.column_letter for c in ws[1] if c.value == "Status"][0]
                for c in ws[st][1:]:
                    c.fill = PatternFill("solid", fgColor=colour.get(c.value, "FFFFFF"))


def main(usage_dir: str):
    usage = pathlib.Path(usage_dir)
    stamp = dt.datetime.now().strftime("%Y-%m-%d_%H%M")
    run = Run(usage / "AI_Cost_Extraction" / stamp)
    az = AzSignIn(run.say)
    for step in (lambda: graph_cost(run), lambda: graph_usage(run), lambda: graph_licences(run, az),
                 lambda: my_roles(run, az), lambda: graph_extra(run), lambda: audit_probe(run, az), lambda: azure(run, az),
                 lambda: power_platform(run, az), lambda: pp_inventory(run, az), lambda: dataverse(run, az),
                 lambda: service_coverage(run), lambda: not_automated(run)):
        try:
            step()
        except (Exception, SystemExit) as e:  # sign-in failure for a whole family: log it and carry on
            run.logf.write(traceback.format_exc())
            run.say(f"[ERROR] step failed: {type(e).__name__}: {str(e)[:300]}")
    xlsx = run.out / f"AI_Cost_Sources_{stamp}.xlsx"
    write_workbook(run, xlsx, stamp)
    run.say(f"\nWorkbook: {xlsx}\nRaw responses: {run.raw}\nLog: {run.out / 'run.log'}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1])
